"""Offline observed-sales extraction; never accepted-order or availability data."""

from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
import hashlib

VERSION = "uci-observed-sales-adapter-v1"
DATASET = "uci-online-retail-352"
SOURCE_URL = "https://archive.ics.uci.edu/dataset/352/online%2Bretail"
ARCHIVE_URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
MEMBER = "Online Retail.xlsx"
ATTRIBUTION = "Chen, D. (2015). Online Retail. UCI Machine Learning Repository. DOI:10.24432/C5BW33."
LICENSE_URL = "https://creativecommons.org/licenses/by/4.0/"
HEADER = ("InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country")
SOURCE_START = date(2010, 12, 1)
SOURCE_LAST_DAY = date(2011, 12, 9)
EXPECTED_ROWS = 541909


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _key(value):
    if type(value) is int and value >= 0:
        return str(value)
    if isinstance(value, str) and value and value.strip() == value:
        return value
    raise ValueError("invalid identity")


def normalize_row(row):
    """Validate required values without rewriting source rows or source clocks."""
    invoice, item = _key(row.get("InvoiceNo")), _key(row.get("StockCode"))
    quantity = row.get("Quantity")
    if type(quantity) is not int:
        raise ValueError("quantity must be a signed whole integer")
    recorded = row.get("InvoiceDate")
    if type(recorded) is datetime and recorded.tzinfo is None:
        day = recorded.date()
    elif type(recorded) is date:
        day = recorded
    else:
        raise ValueError("date must retain a naive source calendar label")
    if not SOURCE_START <= day <= SOURCE_LAST_DAY:
        raise ValueError("date outside declared source range")
    reason = ("cancellation" if invoice.lower().startswith("c") else
              "nonpositive_non_cancelled" if quantity <= 0 else None)
    return dict(invoice=invoice, item=item, day=day, quantity=quantity, exclusion=reason)


def adapt_rows(rows, *, archive_sha256, sheet, expected_rows=EXPECTED_ROWS):
    """Preserve identities and audit all exclusions; any invalidity blocks zeros."""
    if (not isinstance(archive_sha256, str) or len(archive_sha256) != 64
            or any(char not in "0123456789abcdef" for char in archive_sha256)
            or not isinstance(sheet, str) or not sheet
            or type(expected_rows) is not int or expected_rows < 1):
        raise ValueError("invalid archive/sheet/count attestation")
    identities, business_pairs = set(), set()
    counts = Counter()
    invalid = []
    positive = defaultdict(lambda: defaultdict(int))
    seen_items = defaultdict(set)
    first = last = None
    for number, raw in rows:
        counts["raw_rows"] += 1
        if type(number) is not int or number < 2 or number in identities:
            counts["invalid_rows"] += 1
            invalid.append(dict(row_number=number, reason="invalid_or_repeated_source_row_identity"))
            continue
        identities.add(number)
        try:
            parsed = normalize_row(raw)
        except (ValueError, AttributeError) as error:
            counts["invalid_rows"] += 1
            invalid.append(dict(row_number=number, reason=str(error)))
            continue
        first = min(first, parsed["day"]) if first is not None else parsed["day"]
        last = max(last, parsed["day"]) if last is not None else parsed["day"]
        pair = (parsed["invoice"], parsed["item"])
        counts["repeated_invoice_stock_pairs"] += pair in business_pairs
        business_pairs.add(pair)
        seen_items[parsed["item"]].add(parsed["day"])
        if parsed["exclusion"]:
            counts[parsed["exclusion"] + "_rows"] += 1
        else:
            counts["included_rows"] += 1
            counts["gross_positive_units"] += parsed["quantity"]
            positive[parsed["item"]][parsed["day"]] += parsed["quantity"]
    reasons = []
    if counts["raw_rows"] != expected_rows:
        reasons.append("source_row_count_mismatch")
    if counts["invalid_rows"]:
        reasons.append("invalid_required_rows_or_source_identities")
    # A count match alone cannot certify that archive rows were not omitted and
    # replaced with arbitrary different row numbers in a projected input stream.
    if identities != set(range(2, expected_rows + 2)):
        reasons.append("noncontiguous_source_row_identities")
    audit = dict(protocol_version=VERSION, dataset=DATASET, archive_sha256=archive_sha256,
        sheet=sheet, counts={key: counts[key] for key in ("raw_rows", "invalid_rows", "included_rows",
            "cancellation_rows", "nonpositive_non_cancelled_rows", "repeated_invoice_stock_pairs",
            "gross_positive_units")}, invalid=invalid, expected_rows=expected_rows,
        observed_first_day=first.isoformat() if first else None,
        observed_last_day=last.isoformat() if last else None,
        status="not_assessable" if reasons else "assessable", reasons=reasons,
        target="gross positive non-cancelled invoiced units",
        scope="dataset-retailer", availability="unknown", unconstrained_demand="unknown",
        release_assumption="source date end; modeled, not historical ingestion evidence")
    if reasons:
        return dict(audit=audit, series=None, source_seen_days=None)
    n = (SOURCE_LAST_DAY - SOURCE_START).days + 1
    series = {item: [positive[item].get(SOURCE_START + timedelta(days=offset), 0) for offset in range(n)]
              for item in sorted(seen_items)}
    return dict(audit=audit, series=series,
                source_seen_days={item: sorted(day.isoformat() for day in seen_items[item]) for item in sorted(seen_items)})


def read_workbook(path, *, archive_sha256, expected_rows=EXPECTED_ROWS):
    """Optional streaming XLSX reader; formula-required inputs fail validation."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["Online Retail"]:
            raise ValueError("source sheet schema changed")
        sheet = workbook["Online Retail"]
        rows = sheet.iter_rows(values_only=False)
        header = next(rows)
        if tuple(cell.value for cell in header) != HEADER:
            raise ValueError("source column schema changed")
        def projected():
            for number, cells in enumerate(rows, 2):
                values = {name: cell.value for name, cell in zip(HEADER, cells)}
                # A formula string is not an authoritative identity/date/quantity.
                if any(cell.data_type == "f" for name, cell in zip(HEADER, cells)
                       if name in ("InvoiceNo", "StockCode", "Quantity", "InvoiceDate")):
                    values["Quantity"] = None
                yield number, values
        result = adapt_rows(projected(), archive_sha256=archive_sha256, sheet=sheet.title, expected_rows=expected_rows)
        result["audit"]["header"] = list(HEADER)
        result["audit"]["sheet_names"] = workbook.sheetnames
        if result["audit"]["observed_first_day"] != SOURCE_START.isoformat() or result["audit"]["observed_last_day"] != SOURCE_LAST_DAY.isoformat():
            result["audit"]["status"] = "not_assessable"
            result["audit"]["reasons"].append("source_calendar_bounds_changed")
            result["series"] = result["source_seen_days"] = None
        return result
    finally:
        workbook.close()
